# -*- coding: utf-8 -*-
"""Build 'Personalized Product Recommendation' case study as a real .pptx.

Mirrors presentation/recommendation-case-study.html slide-for-slide,
including speaker notes extracted from the HTML.
"""
import html as html_mod
import re
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
HTML_PATH = ROOT / "presentation" / "recommendation-case-study.html"
OUT_PATH = ROOT / "presentation" / "Ayush_Sharma_Recommendation_CaseStudy.pptx"

# ---------- theme ----------
BG = RGBColor.from_string("0B1220")
PANEL = RGBColor.from_string("141D33")
PANEL2 = RGBColor.from_string("101A2E")
LINE = RGBColor.from_string("273349")
CYAN = RGBColor.from_string("22D3EE")
VIOLET = RGBColor.from_string("A78BFA")
AMBER = RGBColor.from_string("FBBF24")
GREEN = RGBColor.from_string("34D399")
INK = RGBColor.from_string("E8EEFB")
BODY = RGBColor.from_string("D7E2F2")
MUTED = RGBColor.from_string("9FB0C8")
CYAN_DARK = RGBColor.from_string("10314A")
VIOLET_DARK = RGBColor.from_string("241C44")
AMBER_DARK = RGBColor.from_string("3A3013")
GREEN_DARK = RGBColor.from_string("12321F")
GRAY = RGBColor.from_string("94A3B8")
WHITE = RGBColor.from_string("FFFFFF")

FONT = "Segoe UI"
W, H = 13.333, 7.5
MX = 0.6              # side margin
CW = W - 2 * MX       # content width

prs = Presentation()
prs.slide_width = Inches(W)
prs.slide_height = Inches(H)
BLANK = prs.slide_layouts[6]


# ---------- low-level helpers ----------
def new_slide():
    s = prs.slides.add_slide(BLANK)
    bg = s.background
    bg.fill.solid()
    bg.fill.fore_color.rgb = BG
    return s


def shape(s, kind, x, y, w, h, fill=PANEL, line=LINE, line_w=0.75, radius=None):
    sp = s.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    sp.shadow.inherit = False
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid()
        sp.fill.fore_color.rgb = fill
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line
        sp.line.width = Pt(line_w)
    if radius is not None and kind == MSO_SHAPE.ROUNDED_RECTANGLE:
        try:
            sp.adjustments[0] = radius
        except Exception:
            pass
    tf = sp.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.08)
    tf.margin_top = tf.margin_bottom = Inches(0.04)
    return sp


def card(s, x, y, w, h, fill=PANEL, line=LINE, radius=0.06):
    return shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, fill=fill, line=line, radius=radius)


def textbox(s, x, y, w, h, anchor=MSO_ANCHOR.TOP):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = 0
    tf.margin_top = tf.margin_bottom = 0
    return tb


def para(tf, first=False, align=PP_ALIGN.LEFT, space_after=0, line_spacing=1.0):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align
    p.space_after = Pt(space_after)
    p.line_spacing = line_spacing
    return p


def run(p, text, size=14, color=BODY, bold=False, italic=False):
    r = p.add_run()
    r.text = text
    r.font.name = FONT
    r.font.size = Pt(size)
    r.font.color.rgb = color
    r.font.bold = bold
    r.font.italic = italic
    return r


def put_lines(sp, lines, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    """lines: list of lists of (text,size,color,bold) run tuples."""
    tf = sp.text_frame
    tf.vertical_anchor = anchor
    for i, runs in enumerate(lines):
        p = para(tf, first=(i == 0), align=align, space_after=3, line_spacing=1.12)
        for t, size, color, bold in runs:
            run(p, t, size=size, color=color, bold=bold)
    return sp


def bullets(s, x, y, w, h, items, size=14, gap=7, marker_color=CYAN):
    """items: str or (lead, rest) tuples."""
    tb = textbox(s, x, y, w, h)
    tf = tb.text_frame
    for i, it in enumerate(items):
        p = para(tf, first=(i == 0), space_after=gap, line_spacing=1.14)
        run(p, "\u25b8  ", size=size, color=marker_color, bold=True)
        if isinstance(it, tuple):
            lead, rest = it
            run(p, lead, size=size, color=WHITE, bold=True)
            if rest:
                run(p, rest, size=size, color=BODY)
        else:
            run(p, it, size=size, color=BODY)
    return tb


def chip(s, x, y, text, color=CYAN, dark=CYAN_DARK, size=10.5):
    w = 0.34 + 0.082 * len(text) * (size / 11.0)
    sp = shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, 0.32, fill=dark, line=color, radius=0.5)
    put_lines(sp, [[(text, size, color, True)]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    return x + w + 0.14


def chip_row(s, x, y, texts_color_specs, size=10.5):
    cx = x
    for text, color, dark in texts_color_specs:
        cx = chip(s, cx, y, text, color=color, dark=dark, size=size)
    return cx


def flow_node(s, x, y, w, h, text, fill=PANEL2, line=LINE, size=12, bold=True):
    sp = shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, fill=fill, line=line, radius=0.18)
    put_lines(sp, [[(text, size, WHITE if bold else BODY, bold)]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    return sp


def arrow_text(s, x, y, w=0.3, h=0.3, glyph="\u2192", color=CYAN, size=14):
    tb = textbox(s, x, y, w, h, anchor=MSO_ANCHOR.MIDDLE)
    p = para(tb.text_frame, first=True, align=PP_ALIGN.CENTER)
    run(p, glyph, size=size, color=color, bold=True)
    return tb


def header(s, kicker, title, title_extra=None):
    shape(s, MSO_SHAPE.RECTANGLE, MX, 0.5, 0.075, 0.78, fill=CYAN, line=None)
    tb = textbox(s, MX + 0.22, 0.42, CW - 0.3, 0.3)
    p = para(tb.text_frame, first=True)
    run(p, kicker.upper(), size=11, color=CYAN, bold=True)
    tb2 = textbox(s, MX + 0.22, 0.66, CW - 0.3, 0.65)
    p2 = para(tb2.text_frame, first=True)
    run(p2, title, size=26, color=WHITE, bold=True)
    if title_extra:
        run(p2, "   " + title_extra[0], size=11.5, color=title_extra[1], bold=True)
    return 1.55  # content top


def footer(s, n, left="AI-Driven Data Analysis & Visualization"):
    tb = textbox(s, MX, H - 0.42, CW, 0.3)
    p = para(tb.text_frame, first=True)
    run(p, left, size=9, color=MUTED)
    tb2 = textbox(s, MX, H - 0.42, CW, 0.3)
    p2 = para(tb2.text_frame, first=True, align=PP_ALIGN.RIGHT)
    run(p2, f"Ayush Sharma \u00b7 B.Tech CSE \u00b7 {n} / 18", size=9, color=MUTED)
    ln = shape(s, MSO_SHAPE.RECTANGLE, MX, H - 0.5, CW, 0.012, fill=LINE, line=None)


def table(s, x, y, w, col_ws, rows, font=11.5, row_h=0.36, header_row=True,
          right_cols=(), hi_rows=(), body_color=BODY):
    n_r, n_c = len(rows), len(col_ws)
    gfx = s.shapes.add_table(n_r, n_c, Inches(x), Inches(y), Inches(w), Inches(row_h * n_r))
    tbl = gfx.table
    tbl.first_row = False
    tbl.horz_banding = False
    for i, cw_ in enumerate(col_ws):
        tbl.columns[i].width = Inches(cw_)
    for r_i, row in enumerate(rows):
        tbl.rows[r_i].height = Inches(row_h)
        for c_i, val in enumerate(row):
            c = tbl.cell(r_i, c_i)
            c.margin_left = c.margin_right = Inches(0.09)
            c.margin_top = c.margin_bottom = Inches(0.03)
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            is_head = header_row and r_i == 0
            if is_head:
                c.fill.solid(); c.fill.fore_color.rgb = CYAN_DARK
                col, bold, size = RGBColor.from_string("E9FBFF"), True, font
            elif r_i in hi_rows:
                c.fill.solid(); c.fill.fore_color.rgb = VIOLET_DARK
                col, bold, size = WHITE, True, font
            else:
                c.fill.solid(); c.fill.fore_color.rgb = PANEL if r_i % 2 else PANEL2
                col, bold, size = body_color, False, font
            tf = c.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.RIGHT if c_i in right_cols and not is_head else PP_ALIGN.LEFT
            run(p, str(val), size=size, color=col, bold=bold)
    return tbl


def set_notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def extract_notes():
    raw = HTML_PATH.read_text(encoding="utf-8")
    blocks = re.findall(r'<div class="notes">(.*?)</div>', raw, flags=re.S)
    return [re.sub(r"\s+", " ", html_mod.unescape(b)).strip() for b in blocks]


def net_art(s):
    """Subtle neural-network decoration for the title slide."""
    nodes = [(10.0, 1.2), (11.2, 2.1), (10.4, 3.1), (11.9, 1.4), (11.5, 4.0), (12.3, 2.8), (10.9, 5.0), (12.5, 4.9)]
    edges = [(0, 1), (1, 2), (0, 3), (1, 5), (2, 4), (4, 5), (2, 6), (4, 7), (5, 7)]
    for a, b in edges:
        x1, y1 = nodes[a]; x2, y2 = nodes[b]
        cn = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
        cn.line.color.rgb = RGBColor.from_string("1E3A52")
        cn.line.width = Pt(1)
        cn.shadow.inherit = False
    for i, (x, y) in enumerate(nodes):
        d = 0.13 if i % 2 else 0.18
        shape(s, MSO_SHAPE.OVAL, x - d / 2, y - d / 2, d, d, fill=CYAN if i % 2 else VIOLET, line=None)


def bar_metric(s, x, y, w, label, val, color):
    tb = textbox(s, x, y, 1.05, 0.24)
    p = para(tb.text_frame, first=True, align=PP_ALIGN.RIGHT)
    run(p, label, size=10, color=MUTED)
    full = w - 1.05 - 0.5
    shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, x + 1.13, y + 0.02, max(full * val, 0.05), 0.19,
          fill=color, line=None, radius=0.5)
    tb2 = textbox(s, x + 1.13 + full + 0.08, y, 0.45, 0.24)
    p2 = para(tb2.text_frame, first=True)
    run(p2, f"{val:.2f}", size=10, color=BODY)


def group_label(s, x, y, text):
    tb = textbox(s, x, y, 4, 0.24)
    p = para(tb.text_frame, first=True)
    run(p, text.upper(), size=10, color=CYAN, bold=True)


# ---------- build slides ----------
NOTES = extract_notes()
assert len(NOTES) == 18, f"expected 18 notes, found {len(NOTES)}"

# ===== 1 · TITLE =====
s = new_slide()
net_art(s)
cx = chip(s, MX, 1.15, "CASE STUDY PRESENTATION", color=AMBER, dark=AMBER_DARK, size=11)
tb = textbox(s, MX, 1.75, 9.2, 2.4)
tf = tb.text_frame
p = para(tf, first=True, line_spacing=1.08)
run(p, "Personalized Product Recommendation", size=40, color=WHITE, bold=True)
p = para(tf, line_spacing=1.08)
run(p, "using Purchase History ", size=40, color=CYAN, bold=True)
run(p, "& Browsing Behavior", size=40, color=VIOLET, bold=True)
shape(s, MSO_SHAPE.RECTANGLE, MX + 0.02, 4.15, 3.4, 0.045, fill=CYAN, line=None)
tb = textbox(s, MX, 4.42, 9.5, 0.4)
p = para(tb.text_frame, first=True)
run(p, "Course: ", size=15, color=MUTED)
run(p, "AI-Driven Data Analysis & Visualization", size=15, color=WHITE, bold=True)
run(p, "  \u00b7  B.Tech CSE \u00b7 Indus University", size=15, color=MUTED)
cd = card(s, MX, 5.25, 5.9, 1.05)
put_lines(cd, [
    [("PRESENTED BY", 10, MUTED, True)],
    [("Ayush Sharma", 18, WHITE, True)],
], anchor=MSO_ANCHOR.MIDDLE)
cd2 = card(s, 6.65, 5.25, 3.6, 1.05)
put_lines(cd2, [
    [("PROGRAM", 10, MUTED, True)],
    [("B.Tech CSE", 14, BODY, False)],
], anchor=MSO_ANCHOR.MIDDLE)
footer(s, 1, "Case Study \u00b7 2026")
set_notes(s, NOTES[0])

# ===== 2 · ABSTRACT =====
s = new_slide()
top = header(s, "02 \u00b7 Abstract", "Abstract")
bullets(s, MX, top + 0.15, 7.35, 4.9, [
    ("Personalized recommendation", " \u2014 AI tailors product suggestions to each user instead of one list for everyone."),
    ("Why these two signals?", " Purchase history captures long-term preference; browsing captures short-term, real-time intent."),
    ("AI processing", " \u2014 both sources are cleaned, converted to numerical features, and fused into one user profile."),
    ("Approach", " \u2014 hybrid recommender combining collaborative filtering, content-based filtering and ML ranking."),
    ("Impact", " \u2014 better discovery, higher engagement and improved conversion for e-commerce platforms."),
], size=14.5, gap=12)
cd = card(s, 8.3, top + 0.2, 4.43, 3.3)
cd.line.color.rgb = CYAN
put_lines(cd, [
    [("In one line", 15, CYAN, True)],
    [("An AI pipeline that learns ", 13.5, BODY, False), ("who the user is", 13.5, WHITE, True),
     (" (purchases) and ", 13.5, BODY, False), ("what they want now", 13.5, WHITE, True),
     (" (browsing), fuses both views, and ranks the most relevant products.", 13.5, BODY, False)],
])
chip_row(s, 8.3, top + 3.75, [("Data", CYAN, CYAN_DARK), ("\u2192", MUTED, None),
                              ("Features", CYAN, CYAN_DARK), ("\u2192", MUTED, None),
                              ("Hybrid Model", VIOLET, VIOLET_DARK), ("\u2192", MUTED, None),
                              ("Top-N", AMBER, AMBER_DARK)], size=10)
footer(s, 2)
set_notes(s, NOTES[1])

# ===== 3 · INTRODUCTION =====
s = new_slide()
top = header(s, "03 \u00b7 Introduction", "Introduction")
bullets(s, MX, top + 0.15, 7.0, 5.0, [
    ("Recommendation system (RS)", " \u2014 filters a huge catalog and ranks the few items a user is most likely to value."),
    ("Personalization in e-commerce", " \u2014 \u201ccustomers who bought this\u2026\u201d shelves, personalized feeds, targeted offers."),
    ("Role of AI/ML", " \u2014 learns patterns from user\u2013item interactions; no hand-written rules for every pair."),
    ("Generic vs personalized", " \u2014 same bestseller list for all users vs a ranking computed per user."),
    ("Two signals beat one", " \u2014 purchases = confirmed preference; browsing = current intent."),
], size=14.5, gap=12)
cd = card(s, 7.95, top + 0.15, 4.78, 3.4)
put_lines(cd, [
    [("Generic vs Personalized", 15, WHITE, True)],
    [("\U0001f310  Generic: ", 12.5, CYAN, True), ("\u201cTop 10 laptops\u201d \u2014 identical for every visitor; ignores taste & budget.", 12.5, BODY, False)],
    [("\U0001f464  Personalized: ", 12.5, VIOLET, True), ("\u201cGaming peripherals for you\u201d \u2014 computed from your purchases + your clicks.", 12.5, BODY, False)],
])
chip_row(s, 7.95, top + 3.75, [("\U0001f6d2 Purchase history", CYAN, CYAN_DARK),
                               ("\U0001f5b1 Browsing behavior", VIOLET, VIOLET_DARK),
                               ("\U0001f9e0 AI model", AMBER, AMBER_DARK)], size=10)
footer(s, 3)
set_notes(s, NOTES[2])

# ===== 4 · PROBLEM STATEMENT =====
s = new_slide()
top = header(s, "04 \u00b7 Problem Statement", "Problem Statement")
bullets(s, MX, top + 0.15, 6.9, 5.0, [
    ("Large catalogs", " \u2014 millions of products; users cannot browse them all."),
    ("Information overload", " \u2014 irrelevant results frustrate users and reduce engagement."),
    ("Different users, different tastes", " \u2014 one ranking cannot fit everyone."),
    ("Purchase history alone is insufficient", " \u2014 sparse, delayed, blind to brand-new interests."),
    ("Browsing adds real-time intent", " \u2014 clicks, searches and dwell time reveal what the user wants now."),
    ("Need", " \u2014 an intelligent system that combines both signals into one ranking."),
], size=14, gap=9)
fx = 7.95
fw = 4.78
n1 = flow_node(s, fx, top + 0.15, fw, 1.0, "", fill=AMBER_DARK, line=AMBER)
put_lines(n1, [[("\u26a0  PROBLEM", 14, AMBER, True)], [("Overload + single-signal recommendations", 11.5, BODY, False)]],
          align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
arrow_text(s, fx + fw / 2 - 0.15, top + 1.22, glyph="\u25bc", color=CYAN, size=16)
n2 = flow_node(s, fx, top + 1.62, fw, 1.0, "", fill=CYAN_DARK, line=CYAN)
put_lines(n2, [[("\U0001f9e0  AI SOLUTION", 14, CYAN, True)], [("Hybrid model fusing purchases + browsing", 11.5, BODY, False)]],
          align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
arrow_text(s, fx + fw / 2 - 0.15, top + 2.69, glyph="\u25bc", color=CYAN, size=16)
n3 = flow_node(s, fx, top + 3.09, fw, 1.0, "", fill=VIOLET_DARK, line=VIOLET)
put_lines(n3, [[("\U0001f3af  EXPECTED OUTCOME", 14, VIOLET, True)], [("Relevant Top-N \u00b7 engagement & conversion", 11.5, BODY, False)]],
          align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
footer(s, 4)
set_notes(s, NOTES[3])

# ===== 5 · TYPES OF DATA =====
s = new_slide()
top = header(s, "05 \u00b7 Data Understanding", "Types of Data Used")
rows = [
    ["Data Type", "Examples", "Characteristics", "Purpose"],
    ["\U0001f6d2  Purchase History", "Product, quantity, price, category", "Historical \u2014 confirmed, low-frequency", "Long-term preferences"],
    ["\U0001f5b1  Browsing Behavior", "Clicks, views, searches, time spent", "Behavioral \u2014 high-frequency, current", "Short-term / real-time intent"],
    ["\U0001f4e6  Product Data", "Category, brand, price, features", "Item metadata \u2014 static per item", "Product similarity"],
    ["\U0001f464  User Data", "User ID, demographics / preferences", "User attributes \u2014 stable", "Personalization"],
]
table(s, MX, top + 0.2, CW, [3.1, 3.4, 3.2, 2.43], rows, font=12.5, row_h=0.72)
chip_row(s, MX, top + 4.15, [("Historical \u2192 who the user has been", AMBER, AMBER_DARK),
                             ("Behavioral \u2192 what the user wants now", CYAN, CYAN_DARK),
                             ("Metadata & attributes \u2192 the glue for similarity", VIOLET, VIOLET_DARK)], size=11)
footer(s, 5)
set_notes(s, NOTES[4])

# ===== 6 · DATA COLLECTION =====
s = new_slide()
top = header(s, "06 \u00b7 Data Collection", "Where the Data Comes From")
bullets(s, MX, top + 0.1, 6.7, 3.0, [
    ("Transaction databases", " \u2014 orders, line items, payment records"),
    ("Clickstream logs", " \u2014 page views, clicks, dwell time"),
    ("Search query logs", " \u2014 what users type into search"),
    ("Product catalog", " \u2014 structured item metadata"),
    ("Interaction / session logs", " \u2014 carts, wishlists, session IDs"),
], size=14, gap=8)
cd = card(s, MX, top + 3.35, 6.7, 1.0, fill=GREEN_DARK, line=GREEN)
put_lines(cd, [
    [("\U0001f512  Privacy-conscious collection", 13, GREEN, True)],
    [("Informed consent \u00b7 anonymization \u00b7 minimal retention \u00b7 GDPR-style compliance", 11.5, BODY, False)],
])
rows = [
    ["Field", "Example"],
    ["user_id", "U-1042"],
    ["product_id", "P-8801"],
    ["timestamp", "2026-03-14 18:22:07"],
    ["action", "view / click / search / purchase"],
    ["quantity", "1"],
    ["category", "Electronics \u203a Peripherals"],
    ["price", "\u20b9 4,999"],
    ["session_id", "S-20260314-771"],
]
cdt = card(s, 7.6, top + 0.1, 5.13, 4.35)
put_lines(cdt, [[("Example interaction records ", 13, WHITE, True), ("(illustrative schema)", 10.5, MUTED, False)]])
table(s, 7.85, top + 0.62, 4.6, [1.5, 3.1], rows, font=10.5, row_h=0.41)
footer(s, 6)
set_notes(s, NOTES[5])

# ===== 7 · PREPROCESSING =====
s = new_slide()
top = header(s, "07 \u00b7 Data Preprocessing", "From Raw Logs to Model-Ready Data")
r1 = ["Raw Data", "Cleaning", "Missing Values", "Deduplication"]
r2 = ["Normalization", "Sessionization", "Feature Engineering", "Model-Ready Data"]
xs = MX
nw = 2.35
for i, t in enumerate(r1):
    hl = t in ("Feature Engineering", "Model-Ready Data")
    flow_node(s, xs, top + 0.15, nw, 0.52, t, fill=CYAN_DARK if t == "Feature Engineering" else (VIOLET_DARK if t == "Model-Ready Data" else PANEL2),
              line=CYAN if t == "Feature Engineering" else (VIOLET if t == "Model-Ready Data" else LINE), size=12)
    if i < 3:
        arrow_text(s, xs + nw + 0.02, top + 0.24)
    xs += nw + 0.34
arrow_text(s, W - MX - 0.45, top + 0.78, glyph="\u2193", color=CYAN, size=15)
xs = MX
for i, t in enumerate(r2):
    flow_node(s, xs, top + 1.1, nw, 0.52, t, fill=CYAN_DARK if t == "Feature Engineering" else (VIOLET_DARK if t == "Model-Ready Data" else PANEL2),
              line=CYAN if t == "Feature Engineering" else (VIOLET if t == "Model-Ready Data" else LINE), size=12)
    if i < 3:
        arrow_text(s, xs + nw + 0.02, top + 1.19)
    xs += nw + 0.34
cd = card(s, MX, top + 2.1, 5.96, 2.75)
put_lines(cd, [[("Common data quality issues", 14, WHITE, True)]])
bullets(s, MX + 0.25, top + 2.62, 5.5, 2.2, [
    ("Missing values", " \u2014 absent price/category \u2192 impute or drop"),
    ("Duplicate events", " \u2014 same click logged twice \u2192 remove"),
    ("Invalid records", " \u2014 bot traffic, refunds, test orders"),
    ("Timestamps", " \u2014 unify time zones; derive time features"),
], size=12, gap=5)
cd2 = card(s, 6.77, top + 2.1, 5.96, 2.75)
put_lines(cd2, [[("Transformations", 14, WHITE, True)]])
bullets(s, 7.02, top + 2.62, 5.5, 2.2, [
    ("Categorical encoding", " \u2014 one-hot / embeddings for category, brand"),
    ("Numerical normalization", " \u2014 comparable scales for price, counts"),
    ("Sessionization", " \u2014 group events with > 30-min inactivity gaps"),
    ("Output", " \u2014 tidy user\u2013item\u2013event tables, model-ready"),
], size=12, gap=5)
footer(s, 7)
set_notes(s, NOTES[6])

# ===== 8 · AI TECHNIQUES =====
s = new_slide()
top = header(s, "08 \u00b7 AI Techniques", "AI Techniques Used")
cards_data = [
    ("1 \u00b7 Collaborative Filtering", "\u201cUsers like you bought\u2026\u201d \u2014 learns from the user\u2013item interaction matrix.", LINE, PANEL),
    ("2 \u00b7 Content-Based Filtering", "Recommends items similar to what the user liked, using product metadata.", LINE, PANEL),
    ("3 \u00b7 Hybrid Recommendation", "Combines both \u2014 long-term taste AND short-term intent. Our choice.", VIOLET, VIOLET_DARK),
    ("4 \u00b7 ML Ranking (LTR)", "Re-orders candidates by predicted relevance \u2014 final Top-N from ranking.", LINE, PANEL),
    ("5 \u00b7 Neural / Deep Learning", "Embedding & wide-and-deep models learn complex interactions (optional).", LINE, PANEL),
    ("Why hybrid?", "CF alone \u2192 cold-start. Content alone \u2192 no serendipity. Purchases \u2192 no current intent. Hybrid exploits all.", VIOLET, VIOLET_DARK),
]
cw_, ch_, gx, gy = 3.91, 1.14, 0.2, 0.16
for i, (t, b, ln, fl) in enumerate(cards_data):
    x = MX + (i % 3) * (cw_ + gx)
    y = top + 0.05 + (i // 3) * (ch_ + gy)
    cd = card(s, x, y, cw_, ch_, fill=fl, line=ln)
    put_lines(cd, [[(t, 12.5, VIOLET if ln == VIOLET else WHITE, True)], [(b, 10, BODY, False)]])
rows = [
    ["Technique", "Purchase data", "Browsing data", "Cold-start", "Best for"],
    ["Collaborative Filtering", "\u2714", "\u2714 (interactions)", "Weak", "Users with rich history"],
    ["Content-Based", "Indirect", "\u2714 (viewed items)", "Good (new items)", "New / niche items"],
    ["Hybrid", "\u2714", "\u2714", "Mitigated", "Accuracy + coverage"],
]
table(s, MX, top + 2.95, CW, [2.9, 1.9, 2.3, 2.2, 2.83], rows, font=11, row_h=0.5, hi_rows=(3,))
footer(s, 8)
set_notes(s, NOTES[7])

# ===== 9 · FEATURE ENGINEERING =====
s = new_slide()
top = header(s, "09 \u00b7 Feature Engineering", "Turning Behavior into Numbers")
cd = card(s, MX, top + 0.1, 5.96, 2.9)
cd.line.color.rgb = CYAN
put_lines(cd, [[("\U0001f6d2  Purchase features \u2014 long-term profile", 14, CYAN, True)]])
bullets(s, MX + 0.25, top + 0.62, 5.5, 2.3, [
    ("Purchase frequency", " \u2014 orders per month"),
    ("Recency", " \u2014 days since last purchase"),
    ("Total spending", " \u2014 lifetime monetary value"),
    ("Category preference", " \u2014 share of spend per category"),
    ("Brand preference", " \u2014 repeat-purchase brand affinity"),
], size=12, gap=4, marker_color=CYAN)
cd2 = card(s, 6.77, top + 0.1, 5.96, 2.9)
cd2.line.color.rgb = VIOLET
put_lines(cd2, [[("\U0001f5b1  Browsing features \u2014 short-term intent", 14, VIOLET, True)]])
bullets(s, 7.02, top + 0.62, 5.5, 2.3, [
    ("Number of views", " \u2014 per item / category"),
    ("Click frequency", " \u2014 within session & recent window"),
    ("Search frequency", " \u2014 queries per session"),
    ("Time spent", " \u2014 dwell time on product pages"),
    ("Recently viewed", " \u2014 last-N items with decay"),
], size=12, gap=4, marker_color=VIOLET)
fl = ["Raw event log", "Aggregate per user / session", "Encode + normalize", "Numerical vector x\u1d64"]
xs = MX
for i, t in enumerate(fl):
    flow_node(s, xs, top + 3.35, 2.55, 0.5, t, size=11,
              fill=VIOLET_DARK if i == 3 else PANEL2, line=VIOLET if i == 3 else LINE)
    if i < 3:
        arrow_text(s, xs + 2.57, top + 3.43, w=0.22, size=13)
    xs += 2.57 + 0.26
tb = textbox(s, MX, top + 4.15, CW, 0.4)
p = para(tb.text_frame, first=True, align=PP_ALIGN.CENTER)
run(p, "Every feature becomes a dimension of the user's vector ", size=12, color=MUTED)
run(p, "x\u1d64", size=12, color=WHITE, bold=True)
run(p, " \u2014 the input the AI model actually learns from.", size=12, color=MUTED)
footer(s, 9)
set_notes(s, NOTES[8])

# ===== 10 · ARCHITECTURE =====
s = new_slide()
top = header(s, "10 \u00b7 System Architecture", "End-to-End Workflow")
col1 = [("USER", PANEL, LINE, True), ("Purchases + Browsing", CYAN_DARK, CYAN, True),
        ("Data Collection Layer", PANEL2, LINE, False), ("Data Preprocessing", PANEL2, LINE, False),
        ("Feature Engineering", PANEL2, LINE, False), ("DATA FUSION", VIOLET_DARK, VIOLET, True)]
col2 = [("HYBRID RECOMMENDATION MODEL", VIOLET_DARK, VIOLET, True), ("Candidate Generation", PANEL2, LINE, False),
        ("Ranking", PANEL2, LINE, False), ("TOP-N RECOMMENDATIONS", CYAN_DARK, CYAN, True),
        ("User Feedback (clicks / purchases)", PANEL2, LINE, False)]
bw, bh, gap = 5.6, 0.5, 0.33
for i, (t, fl, ln, bd) in enumerate(col1):
    flow_node(s, MX, top + i * (bh + gap), bw, bh, t, fill=fl, line=ln, size=12, bold=bd)
    if i < len(col1) - 1:
        arrow_text(s, MX + bw / 2 - 0.15, top + i * (bh + gap) + bh - 0.02, glyph="\u2193", size=14)
for i, (t, fl, ln, bd) in enumerate(col2):
    y = top + i * (bh + gap)
    flow_node(s, 6.95, y, bw, bh, t, fill=fl, line=ln, size=12, bold=bd)
    if i < len(col2) - 1:
        arrow_text(s, 6.95 + bw / 2 - 0.15, y + bh - 0.02, glyph="\u2193", size=14)
loop = card(s, 6.95, top + 5 * (bh + gap) - 0.1, bw, 0.72, fill=PANEL, line=CYAN)
loop.line.dash_style = None
put_lines(loop, [[("\u21ba  Feedback loop \u2014 new interactions retrain the model,", 11, CYAN, True)],
                 [("so it improves continuously", 11, CYAN, True)]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
tb = textbox(s, MX, H - 1.05, CW, 0.45)
p = para(tb.text_frame, first=True, align=PP_ALIGN.CENTER)
run(p, "Two-phase design: ", size=12, color=MUTED)
run(p, "candidate generation", size=12, color=WHITE, bold=True)
run(p, " narrows millions of items to hundreds; ", size=12, color=MUTED)
run(p, "ranking", size=12, color=WHITE, bold=True)
run(p, " orders them precisely.", size=12, color=MUTED)
footer(s, 10)
set_notes(s, NOTES[9])

# ===== 11 · DATA FUSION =====
s = new_slide()
top = header(s, "11 \u00b7 Data Fusion", "Combining the Two Signals")
bullets(s, MX, top + 0.15, 6.6, 4.6, [
    ("Feature-level fusion", " \u2014 concatenate purchase + browsing features into one vector before modeling."),
    ("Score-level fusion", " \u2014 run separate models per source, then combine their scores."),
    ("Weighted hybridization", " \u2014 learn how much each source should contribute."),
    ("User\u2013item interaction matrix", " \u2014 purchases (weight = 1) + views/clicks (fractional) in one matrix."),
    ("Long-term \u2295 short-term", " \u2014 stable profile + fast-changing intent."),
], size=14, gap=10)
cd = card(s, 7.5, top + 0.2, 5.23, 3.6, fill=VIOLET_DARK, line=VIOLET)
put_lines(cd, [
    [("Score fusion model", 15, VIOLET, True)],
    [("Final Score =", 17, INK, True)],
    [("\u03b1\u00b7(Purchase Preference)", 15, AMBER, True)],
    [("+ \u03b2\u00b7(Browsing Interest)", 15, CYAN, True)],
    [("+ \u03b3\u00b7(Product Similarity)", 15, VIOLET, True)],
], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
chip_row(s, 7.5, top + 4.0, [("\u03b1 \u2014 long-term history", AMBER, AMBER_DARK),
                             ("\u03b2 \u2014 current intent", CYAN, CYAN_DARK),
                             ("\u03b3 \u2014 item similarity", VIOLET, VIOLET_DARK)], size=9.5)
tb = textbox(s, 7.5, top + 4.5, 5.23, 0.6)
p = para(tb.text_frame, first=True, align=PP_ALIGN.CENTER)
run(p, "\u03b1, \u03b2, \u03b3 are learned model weights \u2014 tuned on validation data, not hand-picked.", size=11, color=MUTED)
footer(s, 11)
set_notes(s, NOTES[10])

# ===== 12 · TOOLS =====
s = new_slide()
top = header(s, "12 \u00b7 Implementation Stack", "Implementation Tools")
rows = [
    ["Component", "Technology"],
    ["Programming", "Python"],
    ["Data Processing", "Pandas, NumPy"],
    ["Machine Learning", "Scikit-learn"],
    ["Deep Learning", "PyTorch / TensorFlow"],
    ["Database", "PostgreSQL / MongoDB"],
    ["Visualization", "Matplotlib, Seaborn, Plotly"],
    ["API Layer", "FastAPI / Flask"],
    ["Deployment", "Docker / Cloud"],
]
table(s, MX, top + 0.2, 6.4, [2.5, 3.9], rows, font=12.5, row_h=0.55)
cd = card(s, 7.4, top + 0.2, 5.33, 3.6)
cd.line.color.rgb = CYAN
put_lines(cd, [[("Why this stack?", 15, CYAN, True)]])
bullets(s, 7.65, top + 0.75, 4.9, 3.0, [
    ("Python ecosystem", " \u2014 de-facto standard for ML"),
    ("Scikit-learn", " \u2014 fast prototyping of CF + classical models"),
    ("PostgreSQL / MongoDB", " \u2014 relational orders + flexible event logs"),
    ("FastAPI", " \u2014 low-latency REST endpoint for real-time Top-N"),
    ("Docker", " \u2014 reproducible deployment of the pipeline"),
], size=12, gap=7)
chip_row(s, 7.4, top + 4.05, [("Python", CYAN, CYAN_DARK), ("Pandas", CYAN, CYAN_DARK),
                              ("Scikit-learn", VIOLET, VIOLET_DARK), ("FastAPI", AMBER, AMBER_DARK),
                              ("Docker", CYAN, CYAN_DARK)], size=10)
footer(s, 12)
set_notes(s, NOTES[11])

# ===== 13 · WORKED EXAMPLE =====
s = new_slide()
top = header(s, "13 \u00b7 Model in Action", "Worked Example \u2014 How the System Thinks")
cd = card(s, MX, top + 0.1, 5.96, 3.7)
cd.line.color.rgb = CYAN
put_lines(cd, [[("\U0001f464  User A \u2014 observed signals ", 14, CYAN, True), ("(hypothetical)", 10.5, MUTED, False)]])
bullets(s, MX + 0.25, top + 0.62, 5.5, 1.6, [
    ("Purchased:", " Laptop, Wireless Mouse"),
    ("Viewed:", " Gaming Keyboard, Gaming Headset, Monitor"),
    ("Searched:", " \u201cmechanical keyboard\u201d"),
], size=12.5, gap=5)
n1 = flow_node(s, MX + 0.25, top + 2.35, 5.45, 0.55, "Long-term preference \u2192 Computer & accessories",
               fill=CYAN_DARK, line=CYAN, size=11.5)
n2 = flow_node(s, MX + 0.25, top + 3.05, 5.45, 0.55, "Short-term intent \u2192 Gaming / keyboard gear",
               fill=VIOLET_DARK, line=VIOLET, size=11.5)
cd2 = card(s, 6.77, top + 0.1, 5.96, 3.7)
cd2.line.color.rgb = VIOLET
put_lines(cd2, [[("\U0001f3af  Generated Top-N recommendations", 14, VIOLET, True)]])
recs = [("1\ufe0f\u20e3  Mechanical Gaming Keyboard", "search intent + viewed"),
        ("2\ufe0f\u20e3  Gaming Headset", "browsing interest"),
        ("3\ufe0f\u20e3  Gaming Monitor", "complements laptop"),
        ("4\ufe0f\u20e3  Wireless Mouse (upgrade)", "purchase pattern match")]
for i, (t, sub) in enumerate(recs):
    y = top + 0.65 + i * 0.75
    n = flow_node(s, 7.02, y, 5.45, 0.62, "", fill=PANEL2, line=LINE)
    put_lines(n, [[(t, 12, WHITE, True)], [(sub, 9.5, MUTED, False)]], align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE)
tb = textbox(s, 6.77, top + 3.85, 5.96, 0.5)
p = para(tb.text_frame, first=True)
run(p, "\u03b2 (browsing) dominates ranks 1\u20132; \u03b1 (history) shapes ranks 3\u20134 \u2014 fusion weights at work.", size=10.5, color=MUTED)
badge = chip(s, MX, top + 4.35, "\u26a0  HYPOTHETICAL CASE-STUDY EXAMPLE \u2014 ILLUSTRATIVE ONLY, NOT REAL CUSTOMER DATA",
             color=AMBER, dark=AMBER_DARK, size=10.5)
footer(s, 13)
set_notes(s, NOTES[12])

# ===== 14 · RESULTS =====
s = new_slide()
top = header(s, "14 \u00b7 Results & Analysis", "Results & Analysis",
             title_extra=("SIMULATED VALUES \u2014 ACADEMIC DEMONSTRATION ONLY", AMBER))
rows = [
    ["Model", "Precision@5", "Recall@5", "NDCG@5"],
    ["Purchase-only", "0.62", "0.55", "0.60"],
    ["Browsing-only", "0.66", "0.59", "0.64"],
    ["Hybrid (fused)", "0.78", "0.72", "0.76"],
]
table(s, MX, top + 0.15, 6.35, [2.35, 1.4, 1.3, 1.3], rows, font=12, row_h=0.48, right_cols=(1, 2, 3), hi_rows=(3,))
cd = card(s, MX, top + 2.4, 6.35, 1.75, fill=GREEN_DARK, line=GREEN)
put_lines(cd, [[("Metrics used", 13, GREEN, True)],
               [("Precision@K \u2014 share of recommendations that are relevant \u00b7 Recall@K \u2014 share of relevant items captured \u00b7 F1 \u2014 harmonic mean \u00b7 NDCG@K \u2014 rank quality with position discounting \u00b7 CTR / conversion \u2014 online engagement", 10.5, BODY, False)]])
cd2 = card(s, 7.25, top + 0.15, 5.48, 4.0)
put_lines(cd2, [[("Simulated comparison (illustrative)", 13, WHITE, True)]])
groups = [("Precision@5", [(0.62, GRAY), (0.66, CYAN), (0.78, VIOLET)]),
          ("Recall@5", [(0.55, GRAY), (0.59, CYAN), (0.72, VIOLET)]),
          ("NDCG@5", [(0.60, GRAY), (0.64, CYAN), (0.76, VIOLET)])]
labels = ["Purchase", "Browsing", "Hybrid"]
yy = top + 0.6
for gname, vals in groups:
    group_label(s, 7.5, yy, gname)
    yy += 0.28
    for (v, c), lab in zip(vals, labels):
        bar_metric(s, 7.5, yy, 5.0, lab, v, c)
        yy += 0.28
    yy += 0.12
leg_x = 7.5
for lab, c in [("Purchase-only", GRAY), ("Browsing-only", CYAN), ("Hybrid", VIOLET)]:
    shape(s, MSO_SHAPE.RECTANGLE, leg_x, yy + 0.05, 0.16, 0.16, fill=c, line=None)
    tb = textbox(s, leg_x + 0.22, yy + 0.02, 1.5, 0.24)
    p = para(tb.text_frame, first=True)
    run(p, lab, size=10, color=MUTED)
    leg_x += 1.75
tb = textbox(s, MX, top + 4.5, 6.35, 0.8)
p = para(tb.text_frame, first=True, line_spacing=1.15)
run(p, "Why hybrid wins: ", size=12, color=WHITE, bold=True)
run(p, "the two signals have complementary errors \u2014 history is reliable but stale, browsing is current but noisy; fusion combines their evidence.", size=12, color=BODY)
footer(s, 14)
set_notes(s, NOTES[13])

# ===== 15 · ADVANTAGES =====
s = new_slide()
top = header(s, "15 \u00b7 Advantages", "Advantages")
adv = [
    ("\U0001f3af Personalized UX", "Every user sees a feed shaped by their own taste and intent."),
    ("\U0001f50d Better discovery", "Surfaces relevant items the user would never search for."),
    ("\U0001f9ed Less overload", "Millions of products narrowed to a handful of good options."),
    ("\U0001f4c8 Higher engagement", "Relevant suggestions increase clicks, session length, retention."),
    ("\U0001f6d2 Improved conversion", "Intent-aware ranking turns interest into purchases."),
    ("\U0001f9e0 Intent understanding", "Fused signals reveal who the user is and what they want now."),
    ("\u26a1 Real-time adaptation", "Browsing feedback updates recommendations within the session."),
    ("\U0001f4ca Business insights", "Interaction data exposes demand patterns, trends, segments."),
]
cw_, ch_, gx, gy = 2.98, 1.6, 0.14, 0.22
for i, (t, b) in enumerate(adv):
    x = MX + (i % 4) * (cw_ + gx)
    y = top + 0.3 + (i // 4) * (ch_ + gy)
    cd = card(s, x, y, cw_, ch_)
    put_lines(cd, [[(t, 12.5, CYAN, True)], [(b, 10.5, BODY, False)]])
footer(s, 15)
set_notes(s, NOTES[14])

# ===== 16 · CHALLENGES & ETHICS =====
s = new_slide()
top = header(s, "16 \u00b7 Challenges & Ethics", "Challenges & Ethical Issues")
cd = card(s, MX, top + 0.1, 5.96, 3.8)
cd.line.color.rgb = AMBER
put_lines(cd, [[("\u2699\ufe0f  Technical challenges", 14, AMBER, True)]])
bullets(s, MX + 0.25, top + 0.6, 5.5, 3.2, [
    ("Cold-start", " \u2014 new users / items have no data"),
    ("Data sparsity", " \u2014 users touch a tiny fraction of catalog"),
    ("Scalability", " \u2014 scoring millions of items per request"),
    ("Real-time processing", " \u2014 streaming clicks into live recs"),
    ("Model bias", " \u2014 popularity bias hides long-tail items"),
    ("Changing preferences", " \u2014 taste drifts; stale models decay"),
    ("Computational cost", " \u2014 training + serving infrastructure"),
], size=11.5, gap=3, marker_color=AMBER)
cd2 = card(s, 6.77, top + 0.1, 5.96, 3.8)
cd2.line.color.rgb = VIOLET
put_lines(cd2, [[("\u2696\ufe0f  Ethical issues", 14, VIOLET, True)]])
bullets(s, 7.02, top + 0.6, 5.5, 3.2, [
    ("User privacy", " \u2014 behavioral data is personal data"),
    ("Consent", " \u2014 tracking must be informed & revocable"),
    ("Data security", " \u2014 breach of purchase logs is severe"),
    ("Tracking concerns", " \u2014 \u201ccreepiness factor\u201d of following users"),
    ("Algorithmic bias", " \u2014 unfair exposure across sellers"),
    ("Transparency", " \u2014 users deserve to know \u201cwhy this item?\u201d"),
    ("Filter bubbles", " \u2014 over-personalization narrows choice"),
], size=11.5, gap=3, marker_color=VIOLET)
cd3 = card(s, MX, top + 4.1, CW, 0.75, fill=GREEN_DARK, line=GREEN)
put_lines(cd3, [[("Responsible-AI practices:  ", 12, GREEN, True),
                 ("privacy-by-design \u00b7 explicit opt-in consent \u00b7 data minimization & anonymization \u00b7 bias audits \u00b7 explainable recommendations \u00b7 easy opt-out of personalization", 12, BODY, False)]],
          anchor=MSO_ANCHOR.MIDDLE)
footer(s, 16)
set_notes(s, NOTES[15])

# ===== 17 · FUTURE SCOPE + CONCLUSION =====
s = new_slide()
top = header(s, "17 \u00b7 Future Scope & Conclusion", "Future Scope & Conclusion")
tb = textbox(s, MX, top + 0.1, 6.2, 0.35)
p = para(tb.text_frame, first=True)
run(p, "\U0001f680 Future scope", size=15, color=CYAN, bold=True)
fut = [("\u26a1 Real-time recommendations", CYAN), ("\U0001f916 Transformer-based recommenders", VIOLET),
       ("\U0001f3ae Reinforcement learning", CYAN), ("\U0001f4a1 Explainable AI (XAI)", AMBER),
       ("\U0001f510 Federated learning", VIOLET), ("\U0001f6e1\ufe0f Privacy-preserving recommendations", CYAN),
       ("\U0001f5bc\ufe0f Multimodal recommendations", AMBER), ("\U0001f4cd Context-aware recommendations", CYAN)]
fx_, fy_ = MX, top + 0.6
for t, c in fut:
    dark = {"\u26a1": CYAN_DARK}.get(t[0], PANEL2)
    w_ = 0.34 + 0.085 * len(t)
    if fx_ + w_ > MX + 6.4:
        fx_ = MX; fy_ += 0.48
    sp = shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, fx_, fy_, w_, 0.38,
               fill={"C": CYAN_DARK, "V": VIOLET_DARK, "A": AMBER_DARK}[
                   {CYAN: "C", VIOLET: "V", AMBER: "A"}[c]], line=c, radius=0.5)
    put_lines(sp, [[(t, 10.5, c, True)]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    fx_ += w_ + 0.12
cd = card(s, 7.25, top + 0.1, 5.48, 3.4, fill=PANEL, line=GREEN)
put_lines(cd, [[("\U0001f4cc Conclusion", 15, GREEN, True)],
               [("Combining ", 13, BODY, False), ("purchase history", 13, WHITE, True),
                (" and ", 13, BODY, False), ("browsing behavior", 13, WHITE, True),
                (" lets an AI system understand both the user's ", 13, BODY, False),
                ("long-term preferences", 13, WHITE, True), (" and their ", 13, BODY, False),
                ("short-term intent", 13, WHITE, True), (". Through preprocessing, feature engineering and ", 13, BODY, False),
                ("data fusion", 13, WHITE, True), (", a hybrid model produces recommendations more relevant than either signal alone \u2014 with responsible-AI safeguards essential for user trust.", 13, BODY, False)]])
f1 = flow_node(s, MX, top + 4.05, 3.4, 0.5, "Purchases \u2192 who you are", size=11)
plus = textbox(s, 4.08, top + 4.05, 0.4, 0.5, anchor=MSO_ANCHOR.MIDDLE)
p = para(plus.text_frame, first=True, align=PP_ALIGN.CENTER); run(p, "+", size=16, color=CYAN, bold=True)
f2 = flow_node(s, 4.55, top + 4.05, 3.4, 0.5, "Browsing \u2192 what you want now", size=11)
arrow_text(s, 8.03, top + 4.12, w=0.3, size=15)
f3 = flow_node(s, 8.45, top + 4.05, 4.28, 0.5, "Fused understanding \u2192 better recommendations",
               fill=CYAN_DARK, line=CYAN, size=11)
footer(s, 17)
set_notes(s, NOTES[16])

# ===== 18 · REFERENCES =====
s = new_slide()
top = header(s, "18 \u00b7 References", "References", title_extra=("(IEEE format)", MUTED))
refs = [
    "[1] P. Resnick and H. R. Varian, \u201cRecommender systems,\u201d Communications of the ACM, vol. 40, no. 3, pp. 56\u201358, 1997.",
    "[2] G. Adomavicius and A. Tuzhilin, \u201cToward the next generation of recommender systems: A survey of the state-of-the-art and possible extensions,\u201d IEEE Trans. Knowledge and Data Engineering, vol. 17, no. 6, pp. 734\u2013749, 2005.",
    "[3] B. Sarwar, G. Karypis, J. Konstan, and J. Riedl, \u201cItem-based collaborative filtering recommendation algorithms,\u201d in Proc. 10th Int. World Wide Web Conf. (WWW), Hong Kong, 2001, pp. 285\u2013295.",
    "[4] Y. Koren, R. Bell, and C. Volinsky, \u201cMatrix factorization techniques for recommender systems,\u201d Computer, vol. 42, no. 8, pp. 30\u201337, 2009.",
    "[5] R. Burke, \u201cHybrid recommender systems: Survey and experiments,\u201d User Modeling and User-Adapted Interaction, vol. 12, no. 4, pp. 331\u2013370, 2002.",
    "[6] F. Ricci, L. Rokach, and B. Shapira, Eds., Recommender Systems Handbook, 2nd ed. New York, NY, USA: Springer, 2015.",
    "[7] J. Bobadilla, F. Ortega, A. Hernando, and A. Guti\u00e9rrez, \u201cRecommender systems survey,\u201d Knowledge-Based Systems, vol. 46, pp. 109\u2013132, 2013.",
    "[8] P. Cremonesi, Y. Koren, and A. Turrin, \u201cPerformance of recommender algorithms on top-N recommendation tasks,\u201d in Proc. 4th ACM Conf. Recommender Systems (RecSys), Barcelona, Spain, 2010, pp. 39\u201346.",
    "[9] S. Rendle, \u201cFactorization machines,\u201d in Proc. 10th IEEE Int. Conf. Data Mining (ICDM), Sydney, Australia, 2010, pp. 995\u20131000.",
    "[10] H.-T. Cheng et al., \u201cWide & deep learning for recommender systems,\u201d in Proc. 1st Workshop Deep Learning for Recommender Systems (DLRS), Boston, MA, USA, 2016, pp. 7\u201310.",
    "[11] X. Wang, X. He, M. Wang, F. Feng, and T.-S. Chua, \u201cNeural graph collaborative filtering,\u201d in Proc. 42nd Int. ACM SIGIR Conf., Paris, France, 2019, pp. 165\u2013174.",
    "[12] G. Guo, J. Zhang, and N. Yorke-Smith, \u201cA novel Bayesian similarity measure for recommender systems,\u201d in Proc. 23rd Int. Joint Conf. Artificial Intelligence (IJCAI), Beijing, China, 2013, pp. 2619\u20132625.",
]
for i, ref in enumerate(refs):
    col = 0 if i < 6 else 1
    x = MX + col * 6.27
    y = top + 0.15 + (i % 6) * 0.72
    tb = textbox(s, x, y, 6.0, 0.7)
    p = para(tb.text_frame, first=True, line_spacing=1.1)
    num, rest = ref.split("]", 1)
    run(p, num + "]", size=10.5, color=CYAN, bold=True)
    run(p, rest, size=10.5, color=BODY, italic=False)
tb = textbox(s, MX, H - 1.1, CW, 0.4)
p = para(tb.text_frame, first=True)
run(p, "All references are real published works \u2014 verifiable in ACM / IEEE / Springer digital libraries.", size=10.5, color=MUTED, italic=True)
footer(s, 18, left="References")
set_notes(s, NOTES[17])

# ---------- save ----------
prs.core_properties.title = "Personalized Product Recommendation using Purchase History and Browsing Behavior"
prs.core_properties.author = "Ayush Sharma"
prs.core_properties.subject = "AI-Driven Data Analysis & Visualization — Case Study"
OUT_PATH.parent.mkdir(parents=True, exist_ok=True)
prs.save(OUT_PATH)
print("Saved:", OUT_PATH)
print("Slides:", len(prs.slides.__iter__.__self__._sldIdLst))
