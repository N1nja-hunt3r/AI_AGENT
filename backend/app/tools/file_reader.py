"""Async multi-format file reader: pdf, docx, txt, csv, xlsx.

Provides text/table extraction, metadata extraction, and health checks.
"""

from __future__ import annotations

import asyncio
import csv
import logging
import time
from abc import ABC, abstractmethod
from dataclasses import dataclass, field, replace
from pathlib import Path
from typing import Any, Dict, List, Optional, Sequence, Tuple

logger = logging.getLogger(__name__)


# ----------------------------------------------------------------------
# Exceptions
# ----------------------------------------------------------------------

class FileReaderError(Exception):
    """Base error for file reading operations."""


class UnsupportedFileTypeError(FileReaderError):
    """Raised when no reader supports the file's extension."""


class FileReadError(FileReaderError):
    """Raised when a file fails to be opened or parsed."""


class MissingDependencyError(FileReaderError):
    """Raised when an optional parsing dependency is not installed."""


# ----------------------------------------------------------------------
# Data models
# ----------------------------------------------------------------------

@dataclass(frozen=True)
class FileMetadata:
    path: str
    file_name: str
    extension: str
    size_bytes: int
    created_at: float
    modified_at: float
    mime_type: Optional[str] = None
    page_count: Optional[int] = None
    sheet_names: Optional[Tuple[str, ...]] = None
    row_count: Optional[int] = None
    column_count: Optional[int] = None
    extra: Dict[str, Any] = field(default_factory=dict)


@dataclass(frozen=True)
class ExtractedDocument:
    text: str
    metadata: FileMetadata
    pages: Optional[Tuple[str, ...]] = None
    tables: Optional[Tuple[Tuple[Tuple[str, ...], ...], ...]] = None


@dataclass(frozen=True)
class ReaderHealth:
    name: str
    healthy: bool
    available: bool
    latency_ms: float
    detail: Optional[str] = None


@dataclass(frozen=True)
class FileReaderHealthReport:
    healthy: bool
    readers: List[ReaderHealth]
    timestamp: float


def _build_base_metadata(path: Path) -> FileMetadata:
    stat = path.stat()
    return FileMetadata(
        path=str(path),
        file_name=path.name,
        extension=path.suffix.lower().lstrip("."),
        size_bytes=stat.st_size,
        created_at=stat.st_ctime,
        modified_at=stat.st_mtime,
    )


# ----------------------------------------------------------------------
# Base reader
# ----------------------------------------------------------------------

class BaseFileReader(ABC):
    """Abstract base for format-specific file readers."""

    extensions: Tuple[str, ...] = ()

    @property
    def name(self) -> str:
        return type(self).__name__

    @abstractmethod
    def is_available(self) -> bool:
        """Return True if required optional dependencies are installed."""
        ...

    @abstractmethod
    async def read(self, path: Path) -> ExtractedDocument:
        ...

    @abstractmethod
    async def extract_metadata(self, path: Path) -> FileMetadata:
        ...

    async def health_check(self) -> ReaderHealth:
        start = time.monotonic()
        try:
            available = self.is_available()
            detail = None if available else "Required dependency is not installed."
            latency_ms = (time.monotonic() - start) * 1000.0
            return ReaderHealth(
                name=self.name, healthy=available, available=available,
                latency_ms=latency_ms, detail=detail,
            )
        except Exception as exc:  # noqa: BLE001 - degrade gracefully into a health result
            latency_ms = (time.monotonic() - start) * 1000.0
            return ReaderHealth(
                name=self.name, healthy=False, available=False,
                latency_ms=latency_ms, detail=str(exc),
            )


# ----------------------------------------------------------------------
# TXT
# ----------------------------------------------------------------------

class TxtReader(BaseFileReader):
    extensions = (".txt",)

    def is_available(self) -> bool:
        return True

    async def read(self, path: Path) -> ExtractedDocument:
        return await asyncio.to_thread(self._read_sync, path)

    async def extract_metadata(self, path: Path) -> FileMetadata:
        return await asyncio.to_thread(self._metadata_sync, path)

    def _read_sync(self, path: Path) -> ExtractedDocument:
        base_meta = _build_base_metadata(path)
        try:
            text = path.read_text(encoding="utf-8", errors="replace")
        except OSError as exc:
            raise FileReadError(f"Failed to read text file '{path}': {exc}") from exc
        line_count = text.count("\n") + (1 if text and not text.endswith("\n") else 0)
        metadata = replace(base_meta, mime_type="text/plain", row_count=line_count)
        return ExtractedDocument(text=text, metadata=metadata)

    def _metadata_sync(self, path: Path) -> FileMetadata:
        base_meta = _build_base_metadata(path)
        try:
            with path.open("r", encoding="utf-8", errors="replace") as handle:
                line_count = sum(1 for _ in handle)
        except OSError as exc:
            raise FileReadError(f"Failed to read text file '{path}': {exc}") from exc
        return replace(base_meta, mime_type="text/plain", row_count=line_count)


# ----------------------------------------------------------------------
# CSV
# ----------------------------------------------------------------------

class CsvReader(BaseFileReader):
    extensions = (".csv",)

    def is_available(self) -> bool:
        return True

    async def read(self, path: Path) -> ExtractedDocument:
        return await asyncio.to_thread(self._read_sync, path)

    async def extract_metadata(self, path: Path) -> FileMetadata:
        return await asyncio.to_thread(self._metadata_sync, path)

    def _read_sync(self, path: Path) -> ExtractedDocument:
        base_meta = _build_base_metadata(path)
        try:
            with path.open("r", encoding="utf-8", newline="", errors="replace") as handle:
                rows = [tuple(row) for row in csv.reader(handle)]
        except (OSError, csv.Error) as exc:
            raise FileReadError(f"Failed to read CSV file '{path}': {exc}") from exc

        column_count = max((len(row) for row in rows), default=0)
        text = "\n".join(",".join(row) for row in rows)
        metadata = replace(
            base_meta, mime_type="text/csv", row_count=len(rows), column_count=column_count
        )
        tables = (tuple(rows),) if rows else None
        return ExtractedDocument(text=text, metadata=metadata, tables=tables)

    def _metadata_sync(self, path: Path) -> FileMetadata:
        base_meta = _build_base_metadata(path)
        try:
            with path.open("r", encoding="utf-8", newline="", errors="replace") as handle:
                row_count = 0
                column_count = 0
                for row in csv.reader(handle):
                    row_count += 1
                    column_count = max(column_count, len(row))
        except (OSError, csv.Error) as exc:
            raise FileReadError(f"Failed to read CSV file '{path}': {exc}") from exc
        return replace(
            base_meta, mime_type="text/csv", row_count=row_count, column_count=column_count
        )


# ----------------------------------------------------------------------
# PDF
# ----------------------------------------------------------------------

class PdfReader(BaseFileReader):
    extensions = (".pdf",)

    def is_available(self) -> bool:
        try:
            import pypdf  # noqa: F401
        except ImportError:
            return False
        return True

    async def read(self, path: Path) -> ExtractedDocument:
        return await asyncio.to_thread(self._read_sync, path)

    async def extract_metadata(self, path: Path) -> FileMetadata:
        return await asyncio.to_thread(self._metadata_sync, path)

    def _open(self, path: Path) -> Any:
        try:
            import pypdf
        except ImportError as exc:
            raise MissingDependencyError(
                "pypdf is required to read PDF files. Install with `pip install pypdf`."
            ) from exc
        try:
            return pypdf.PdfReader(str(path))
        except Exception as exc:  # noqa: BLE001 - normalize to FileReaderError
            raise FileReadError(f"Failed to open PDF file '{path}': {exc}") from exc

    @staticmethod
    def _doc_info_extra(reader: Any) -> Dict[str, Any]:
        info = reader.metadata
        if not info:
            return {}
        return {str(key).lstrip("/"): str(value) for key, value in dict(info).items()}

    def _read_sync(self, path: Path) -> ExtractedDocument:
        base_meta = _build_base_metadata(path)
        reader = self._open(path)
        pages: List[str] = []
        for page in reader.pages:
            try:
                pages.append(page.extract_text() or "")
            except Exception as exc:  # noqa: BLE001 - skip unreadable page
                logger.warning("Failed to extract text from a page in '%s': %s", path, exc)
                pages.append("")
        metadata = replace(
            base_meta,
            mime_type="application/pdf",
            page_count=len(pages),
            extra=self._doc_info_extra(reader),
        )
        return ExtractedDocument(text="\n\n".join(pages), metadata=metadata, pages=tuple(pages))

    def _metadata_sync(self, path: Path) -> FileMetadata:
        base_meta = _build_base_metadata(path)
        reader = self._open(path)
        try:
            page_count = len(reader.pages)
        except Exception as exc:  # noqa: BLE001 - normalize to FileReaderError
            raise FileReadError(f"Failed to inspect PDF file '{path}': {exc}") from exc
        return replace(
            base_meta,
            mime_type="application/pdf",
            page_count=page_count,
            extra=self._doc_info_extra(reader),
        )


# ----------------------------------------------------------------------
# DOCX
# ----------------------------------------------------------------------

class DocxReader(BaseFileReader):
    extensions = (".docx",)

    def is_available(self) -> bool:
        try:
            import docx  # noqa: F401
        except ImportError:
            return False
        return True

    async def read(self, path: Path) -> ExtractedDocument:
        return await asyncio.to_thread(self._read_sync, path)

    async def extract_metadata(self, path: Path) -> FileMetadata:
        return await asyncio.to_thread(self._metadata_sync, path)

    def _open(self, path: Path) -> Any:
        try:
            import docx
        except ImportError as exc:
            raise MissingDependencyError(
                "python-docx is required to read DOCX files. Install with `pip install python-docx`."
            ) from exc
        try:
            return docx.Document(str(path))
        except Exception as exc:  # noqa: BLE001 - normalize to FileReaderError
            raise FileReadError(f"Failed to open DOCX file '{path}': {exc}") from exc

    @staticmethod
    def _core_props_extra(document: Any) -> Dict[str, Any]:
        props = document.core_properties
        return {
            "author": props.author or "",
            "title": props.title or "",
            "subject": props.subject or "",
            "created": props.created.isoformat() if props.created else "",
            "modified": props.modified.isoformat() if props.modified else "",
        }

    def _read_sync(self, path: Path) -> ExtractedDocument:
        base_meta = _build_base_metadata(path)
        document = self._open(path)
        paragraphs = [p.text for p in document.paragraphs]
        tables = tuple(
            tuple(tuple(cell.text for cell in row.cells) for row in table.rows)
            for table in document.tables
        )
        metadata = replace(
            base_meta,
            mime_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            row_count=len(paragraphs),
            extra=self._core_props_extra(document),
        )
        return ExtractedDocument(
            text="\n".join(paragraphs),
            metadata=metadata,
            tables=tables or None,
        )

    def _metadata_sync(self, path: Path) -> FileMetadata:
        base_meta = _build_base_metadata(path)
        document = self._open(path)
        return replace(
            base_meta,
            mime_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            row_count=len(document.paragraphs),
            extra=self._core_props_extra(document),
        )


# ----------------------------------------------------------------------
# XLSX
# ----------------------------------------------------------------------

class XlsxReader(BaseFileReader):
    extensions = (".xlsx",)

    def is_available(self) -> bool:
        try:
            import openpyxl  # noqa: F401
        except ImportError:
            return False
        return True

    async def read(self, path: Path) -> ExtractedDocument:
        return await asyncio.to_thread(self._read_sync, path)

    async def extract_metadata(self, path: Path) -> FileMetadata:
        return await asyncio.to_thread(self._metadata_sync, path)

    def _load(self, path: Path, *, read_only: bool) -> Any:
        try:
            import openpyxl
        except ImportError as exc:
            raise MissingDependencyError(
                "openpyxl is required to read XLSX files. Install with `pip install openpyxl`."
            ) from exc
        try:
            return openpyxl.load_workbook(str(path), data_only=True, read_only=read_only)
        except Exception as exc:  # noqa: BLE001 - normalize to FileReaderError
            raise FileReadError(f"Failed to open XLSX file '{path}': {exc}") from exc

    def _read_sync(self, path: Path) -> ExtractedDocument:
        base_meta = _build_base_metadata(path)
        workbook = self._load(path, read_only=True)
        try:
            sheet_names = tuple(workbook.sheetnames)
            tables: List[Tuple[Tuple[str, ...], ...]] = []
            text_parts: List[str] = []
            total_rows = 0
            max_cols = 0

            for sheet_name in sheet_names:
                sheet = workbook[sheet_name]
                rows: List[Tuple[str, ...]] = []
                for row in sheet.iter_rows(values_only=True):
                    str_row = tuple("" if value is None else str(value) for value in row)
                    rows.append(str_row)
                    max_cols = max(max_cols, len(str_row))
                total_rows += len(rows)
                tables.append(tuple(rows))
                text_parts.append(f"# {sheet_name}\n" + "\n".join(",".join(r) for r in rows))
        finally:
            workbook.close()

        metadata = replace(
            base_meta,
            mime_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            sheet_names=sheet_names,
            row_count=total_rows,
            column_count=max_cols,
        )
        return ExtractedDocument(
            text="\n\n".join(text_parts), metadata=metadata, tables=tuple(tables)
        )

    def _metadata_sync(self, path: Path) -> FileMetadata:
        base_meta = _build_base_metadata(path)
        workbook = self._load(path, read_only=True)
        try:
            sheet_names = tuple(workbook.sheetnames)
            total_rows = 0
            max_cols = 0
            for sheet_name in sheet_names:
                sheet = workbook[sheet_name]
                total_rows += sheet.max_row or 0
                max_cols = max(max_cols, sheet.max_column or 0)
        finally:
            workbook.close()

        return replace(
            base_meta,
            mime_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            sheet_names=sheet_names,
            row_count=total_rows,
            column_count=max_cols,
        )


# ----------------------------------------------------------------------
# Service
# ----------------------------------------------------------------------

def _default_readers() -> Tuple[BaseFileReader, ...]:
    return (TxtReader(), CsvReader(), PdfReader(), DocxReader(), XlsxReader())


class FileReaderService:
    """Dispatches reading and metadata extraction to format-specific readers."""

    def __init__(self, readers: Optional[Sequence[BaseFileReader]] = None) -> None:
        self._readers_by_ext: Dict[str, BaseFileReader] = {}
        for reader in readers or _default_readers():
            for extension in reader.extensions:
                self._readers_by_ext[extension.lower().lstrip(".")] = reader

    def _resolve_reader(self, path: Path) -> BaseFileReader:
        extension = path.suffix.lower().lstrip(".")
        reader = self._readers_by_ext.get(extension)
        if reader is None:
            raise UnsupportedFileTypeError(f"No reader registered for extension '.{extension}'.")
        return reader

    @staticmethod
    def _resolve_path(file_path: str) -> Path:
        path = Path(file_path)
        if not path.exists():
            raise FileNotFoundError(f"File not found: {file_path}")
        if not path.is_file():
            raise FileReadError(f"Path is not a regular file: {file_path}")
        return path

    async def read(self, file_path: str) -> ExtractedDocument:
        path = self._resolve_path(file_path)
        reader = self._resolve_reader(path)
        return await reader.read(path)

    async def extract_metadata(self, file_path: str) -> FileMetadata:
        path = self._resolve_path(file_path)
        reader = self._resolve_reader(path)
        return await reader.extract_metadata(path)

    async def health_check(self) -> FileReaderHealthReport:
        unique_readers = list({id(reader): reader for reader in self._readers_by_ext.values()}.values())
        results = list(await asyncio.gather(*(reader.health_check() for reader in unique_readers)))
        overall_healthy = all(result.healthy for result in results) if results else True
        return FileReaderHealthReport(
            healthy=overall_healthy, readers=results, timestamp=time.time()
        )
